import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import (
    Font,
    PatternFill,
    Alignment,
    Border,
    Side
)
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font
from openpyxl.cell.cell import Hyperlink


HEADER_FILL = PatternFill(
    fill_type="solid",
    start_color="1F4E78"
)

HEADER_FONT = Font(
    color="FFFFFF",
    bold=True
)

ODD_ROW_FILL = PatternFill(
    fill_type="solid",
    start_color="FFFFFF"
)

EVEN_ROW_FILL = PatternFill(
    fill_type="solid",
    start_color="F4F8FC"
)

GREEN_FILL = PatternFill(
    fill_type="solid",
    start_color="C6EFCE"
)

RED_FILL = PatternFill(
    fill_type="solid",
    start_color="FFC7CE"
)

YELLOW_FILL = PatternFill(
    fill_type="solid",
    start_color="FFF2CC"
)

THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin")
)


def format_excel(filename):

    wb = load_workbook(filename)
    ws = wb.active

    # ---------------- Header ----------------

    for cell in ws[1]:

        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = THIN_BORDER

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    # ---------------- Body ----------------

    for row in ws.iter_rows(min_row=2):

        row_fill = (
            EVEN_ROW_FILL
            if row[0].row % 2 == 0
            else ODD_ROW_FILL
        )

        for cell in row:

            cell.fill = row_fill
            cell.border = THIN_BORDER

            cell.alignment = Alignment(
                wrap_text=True,
                vertical="top"
            )

            if cell.value == "Not Available":
                cell.fill = RED_FILL

            if cell.value == "Yes":
                cell.fill = RED_FILL

            if cell.value == "No":
                cell.fill = GREEN_FILL

        # ---------------- Hyperlinks ----------------

    hyperlink_columns = [
        "source_url",
        "composition_source",
        "dosage_source",
        "indication_source"
    ]

    headers = {}

    for cell in ws[1]:
        headers[cell.value] = cell.column

    hyperlink_font = Font(
        color="0563C1",
        underline="single"
    )

    for column_name in hyperlink_columns:

        if column_name not in headers:
            continue

        column_index = headers[column_name]

        for row in range(2, ws.max_row + 1):

            cell = ws.cell(row=row, column=column_index)

            if (
                cell.value
                and str(cell.value).startswith("http")
            ):

                cell.hyperlink = str(cell.value)
                cell.font = hyperlink_font
                
    # ---------------- Auto Width ----------------

    for column in ws.columns:

        max_length = 0

        letter = get_column_letter(column[0].column)

        for cell in column:

            try:
                max_length = max(
                    max_length,
                    len(str(cell.value))
                )
            except:
                pass

        ws.column_dimensions[letter].width = min(
            max_length + 4,
            55
        )

    # ---------------- Freeze ----------------

    ws.freeze_panes = "A2"

    # ---------------- Filter ----------------

    ws.auto_filter.ref = ws.dimensions

    wb.save(filename)


def generate_excel(records, filename):

    df = pd.DataFrame(records)

    df.to_excel(
        filename,
        index=False
    )

    format_excel(filename)


def save_dataframe(df, filename):

    df.to_excel(
        filename,
        index=False
    )

    format_excel(filename)